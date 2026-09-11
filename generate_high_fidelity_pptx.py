import http.server
import socketserver
import threading
import subprocess
import os
import re
import shutil
import time
from pptx import Presentation
from pptx.util import Inches

PROJECT_DIR = r"F:\Aathi\mindbridge"
TEMP_DIR = os.path.join(PROJECT_DIR, "temp_slide_renders")
OUTPUT_PPTX = os.path.join(PROJECT_DIR, "MindBridge_Deck.pptx")
CHROME_PATH = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
PORT = 8899

def run():
    print("=== Generating High-Fidelity Styled PowerPoint Deck ===")
    
    # 1. Create temporary directory
    if os.path.exists(TEMP_DIR):
        shutil.rmtree(TEMP_DIR)
    os.makedirs(TEMP_DIR, exist_ok=True)

    # 2. Extract slides from decks.html
    html_path = os.path.join(PROJECT_DIR, "decks.html")
    with open(html_path, "r", encoding="utf-8") as f:
        full_html = f.read()

    slides = re.findall(r'(<section class="slide-container" id="slide-\d+">.*?</section>)', full_html, re.DOTALL)
    print(f"Extracted {len(slides)} slides from decks.html")

    if len(slides) == 0:
        raise RuntimeError("No slides found in decks.html!")

    # 3. Create standalone HTML for each slide
    slide_html_template = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin="">
<link href="https://fonts.googleapis.com/css2?family=Poppins:wght@400;500;600;700;800;900&display=swap" rel="stylesheet">
<script src="https://cdn.tailwindcss.com"></script>
<style>
  * {{ box-sizing: border-box; }}
  html, body {{
    margin: 0;
    padding: 0;
    width: 1920px;
    height: 1080px;
    overflow: hidden;
    background: #0B0F19;
    font-family: 'Poppins', sans-serif;
  }}
  .slide-container {{
    width: 1920px;
    height: 1080px;
    background-color: #F0F9F8;
    position: relative;
    padding: 80px 100px;
    display: flex;
    flex-direction: column;
    justify-content: space-between;
    overflow: hidden;
  }}
  .gradient-text {{
    background: linear-gradient(135deg, #14B8A6 0%, #38BDF8 100%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
  }}
  .gradient-bg {{
    background: linear-gradient(135deg, #14B8A6 0%, #38BDF8 100%);
  }}
  .gradient-border {{
    border-color: #14B8A6;
  }}
  .heading-xl {{
    font-size: 80px;
    font-weight: 900;
    line-height: 1.1;
    color: #0F172A;
    letter-spacing: -0.02em;
  }}
  .stat-val {{
    font-size: 58px;
    font-weight: 800;
    line-height: 1.1;
  }}
  .label-lg {{
    font-size: 24px;
    font-weight: 400;
    line-height: 1.4;
  }}
</style>
</head>
<body>
{content}
</body>
</html>"""

    for i, slide_content in enumerate(slides):
        slide_file = os.path.join(TEMP_DIR, f"slide_{i+1}.html")
        with open(slide_file, "w", encoding="utf-8") as f:
            f.write(slide_html_template.format(content=slide_content))

    # Copy logo.jpeg into TEMP_DIR so relative image paths work seamlessly
    logo_src = os.path.join(PROJECT_DIR, "logo.jpeg")
    if os.path.exists(logo_src):
        shutil.copy2(logo_src, os.path.join(TEMP_DIR, "logo.jpeg"))

    # 4. Start local HTTP server for rendering
    class Handler(http.server.SimpleHTTPRequestHandler):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, directory=TEMP_DIR, **kwargs)
        def log_message(self, format, *args):
            pass # suppress noisy server logs

    httpd = socketserver.TCPServer(("", PORT), Handler)
    server_thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    server_thread.start()
    print(f"Rendering HTTP server started on port {PORT}")

    # 5. Capture each slide using Chrome headless
    slide_images = []
    for i in range(len(slides)):
        slide_num = i + 1
        img_path = os.path.join(TEMP_DIR, f"slide_{slide_num}.png")
        url = f"http://localhost:{PORT}/slide_{slide_num}.html"
        
        cmd = [
            CHROME_PATH,
            "--headless=new",
            "--disable-gpu",
            "--window-size=1920,1080",
            f"--screenshot={img_path}",
            url
        ]
        print(f"Rendering Slide {slide_num}/{len(slides)}...")
        subprocess.run(cmd, check=True)
        if os.path.exists(img_path):
            slide_images.append(img_path)
        else:
            raise RuntimeError(f"Failed to generate screenshot for slide {slide_num}")

    httpd.shutdown()
    print("All slides rendered successfully!")

    # 6. Build the PowerPoint deck using python-pptx
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen width
    prs.slide_height = Inches(7.5)    # 16:9 widescreen height
    blank_layout = prs.slide_layouts[6]

    slide_titles = [
        "Title: MindBridge Stepped-Care Platform",
        "The Problem: Epidemic in Higher Education",
        "Root Causes: Three Systemic Bottlenecks",
        "Our Solution: MindBridge Architecture & Zero-PII",
        "How It Works: Smart Screening & Dynamic Triage",
        "System Architecture: Pipeline & Privacy Guarantees",
        "Policy Ask: MoE, UGC & MoHFW Integration",
        "Feasibility & Scalability: 90-Day Timeline & Budget",
        "What's Different: Defensibility Comparison Matrix",
        "Impact Potential & Thank You"
    ]

    for i, img_path in enumerate(slide_images):
        slide = prs.slides.add_slide(blank_layout)
        # Place the 1920x1080 image full-bleed
        slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
        
        # Add slide note for presenter
        if i < len(slide_titles):
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = f"Slide {i+1}: {slide_titles[i]}\nMindBridge — YUVA Future 6.0 Health Track Presentation"

    prs.save(OUTPUT_PPTX)
    print(f"SUCCESS: Saved high-fidelity presentation to: {OUTPUT_PPTX}")
    print(f"File size: {os.path.getsize(OUTPUT_PPTX):,} bytes")

    # 7. Clean up temporary render folder
    try:
        shutil.rmtree(TEMP_DIR)
        print("Cleaned up temporary render files.")
    except Exception as e:
        print("Warning: Could not remove temp folder:", e)

if __name__ == "__main__":
    run()
